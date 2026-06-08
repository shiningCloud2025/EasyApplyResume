"""
用 PowerPoint COM 原生生成的动画 XML 结构，一次性生成完整 PPT
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree
import zipfile, uuid

P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

etree.register_namespace('a', A)
etree.register_namespace('r', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships')
etree.register_namespace('p', P)

BG=CARD=RGBColor(0xFF,0xFF,0xFF); BLUE=RGBColor(0x02,0x84,0xC7)
TEAL=RGBColor(0x0D,0x94,0x8F); ORANGE=RGBColor(0xEA,0x58,0x0C)
PURPLE=RGBColor(0x7C,0x3A,0xED); DARK=RGBColor(0x0F,0x17,0x2A)
TEXT=RGBColor(0x1E,0x29,0x3B); GRAY=RGBColor(0x64,0x74,0x8B)
WHITE=RGBColor(0xFF,0xFF,0xFF); LIGHT=RGBColor(0xF1,0xF5,0xF9)

prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
SW=prs.slide_width;SH=prs.slide_height

def bg(s,c):s.background.fill.solid();s.background.fill.fore_color.rgb=c
def rct(s,l,t,w,h,fl,li=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,l,t,w,h)
    sh.fill.solid();sh.fill.fore_color.rgb=fl
    if li:sh.line.fill.solid();sh.line.fill.fore_color.rgb=li
    else:sh.line.fill.background()
    return sh
def rctr(s,l,t,w,h,fl):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,l,t,w,h)
    sh.fill.solid();sh.fill.fore_color.rgb=fl;sh.line.fill.background()
    return sh
def tx(s,l,t,w,h,text,sz=12,color=TEXT,b=False,font='Calibri',al=PP_ALIGN.LEFT,va=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(l,t,w,h);tf=tb.text_frame;tf.word_wrap=True
    p=tf.paragraphs[0];p.text=text;p.font.size=Pt(sz)
    p.font.color.rgb=color;p.font.bold=b;p.font.name=font;p.alignment=al
    return tb
def tbar(s,tt):
    rct(s,Inches(0),Inches(0),SW,Inches(0.8),DARK)
    tx(s,Inches(0.6),Inches(0.15),Inches(12),Inches(0.5),tt,sz=22,color=WHITE,b=True,font='Trebuchet MS',va=MSO_ANCHOR.MIDDLE)
    rct(s,Inches(0),Inches(0.8),SW,Inches(0.03),BLUE)
def pn(s,n):
    tx(s,Inches(11.5),Inches(7.1),Inches(1.5),Inches(0.3),f"{n}/9",sz=9,color=GRAY,al=PP_ALIGN.RIGHT)

def make_timing(shape_ids):
    """用 PowerPoint 原生结构生成 timing XML"""
    t_id = 1
    timing = etree.Element(f'{{{P}}}timing')
    tnLst = etree.SubElement(timing, f'{{{P}}}tnLst')
    root_par = etree.SubElement(tnLst, f'{{{P}}}par')
    root_ctn = etree.SubElement(root_par, f'{{{P}}}cTn', id=str(t_id), dur='indefinite', restart='never', nodeType='tmRoot')
    t_id+=1
    root_cl = etree.SubElement(root_ctn, f'{{{P}}}childTnLst')
    seq = etree.SubElement(root_cl, f'{{{P}}}seq', concurrent='1', nextAc='seek')
    seq_ctn = etree.SubElement(seq, f'{{{P}}}cTn', id=str(t_id), dur='indefinite', nodeType='mainSeq')
    t_id+=1
    seq_cl = etree.SubElement(seq_ctn, f'{{{P}}}childTnLst')
    wrapper_par = etree.SubElement(seq_cl, f'{{{P}}}par')
    wrapper_ctn = etree.SubElement(wrapper_par, f'{{{P}}}cTn', id=str(t_id), fill='hold')
    t_id+=1
    wrapper_sc = etree.SubElement(wrapper_ctn, f'{{{P}}}stCondLst')
    etree.SubElement(wrapper_sc, f'{{{P}}}cond', delay='indefinite')
    wrapper_cl = etree.SubElement(wrapper_ctn, f'{{{P}}}childTnLst')
    inner_par = etree.SubElement(wrapper_cl, f'{{{P}}}par')
    inner_ctn = etree.SubElement(inner_par, f'{{{P}}}cTn', id=str(t_id), fill='hold')
    t_id+=1
    inner_sc = etree.SubElement(inner_ctn, f'{{{P}}}stCondLst')
    etree.SubElement(inner_sc, f'{{{P}}}cond', delay='0')
    inner_cl = etree.SubElement(inner_ctn, f'{{{P}}}childTnLst')
    
    for i, spid in enumerate(shape_ids):
        click_par = etree.SubElement(inner_cl, f'{{{P}}}par')
        click_ctn = etree.SubElement(click_par, f'{{{P}}}cTn', 
            id=str(t_id), presetID='1', presetClass='entr', presetSubtype='0',
            fill='hold', grpId='0', nodeType='clickEffect')
        t_id+=1
        click_sc = etree.SubElement(click_ctn, f'{{{P}}}stCondLst')
        etree.SubElement(click_sc, f'{{{P}}}cond', delay='0')
        click_cl = etree.SubElement(click_ctn, f'{{{P}}}childTnLst')
        set_e = etree.SubElement(click_cl, f'{{{P}}}set')
        cb = etree.SubElement(set_e, f'{{{P}}}cBhvr')
        cb_ctn = etree.SubElement(cb, f'{{{P}}}cTn', id=str(t_id), dur='1', fill='hold')
        t_id+=1
        cb_sc = etree.SubElement(cb_ctn, f'{{{P}}}stCondLst')
        etree.SubElement(cb_sc, f'{{{P}}}cond', delay='0')
        tgt = etree.SubElement(cb, f'{{{P}}}tgtEl')
        etree.SubElement(tgt, f'{{{P}}}spTgt', spid=str(spid))
        anl = etree.SubElement(set_e, f'{{{P}}}attrNameLst')
        etree.SubElement(anl, f'{{{P}}}attrName').text='style.visibility'
        to_e = etree.SubElement(set_e, f'{{{P}}}to')
        etree.SubElement(to_e, f'{{{P}}}strVal', val='visible')
    
    # prev/next cond at seq level
    prev = etree.SubElement(seq, f'{{{P}}}prevCondLst')
    pc = etree.SubElement(prev, f'{{{P}}}cond', evt='onPrev', delay='0')
    pt = etree.SubElement(pc, f'{{{P}}}tgtEl')
    etree.SubElement(pt, f'{{{P}}}sldTgt')
    nxt = etree.SubElement(seq, f'{{{P}}}nextCondLst')
    nc = etree.SubElement(nxt, f'{{{P}}}cond', evt='onNext', delay='0')
    nt = etree.SubElement(nc, f'{{{P}}}tgtEl')
    etree.SubElement(nt, f'{{{P}}}sldTgt')
    
    return timing

def inject_all(slide, trans='fade'):
    """往 slide 的 ZIP XML 层面注入 transition + 原生结构 timing"""
    sld = slide.element
    clr = sld.find(f'{{{P}}}clrMapOvr')
    if clr is not None: sld.remove(clr)
    
    # transition
    te = etree.SubElement(sld, f'{{{P}}}transition', spd='med')
    tmap={'push':'l','cover':'r','wipe':'l','uncover':'d'}
    if trans in tmap:
        t=etree.SubElement(te, f'{{{P}}}{trans}');t.set('dir',tmap[trans])
    elif trans=='dissolve':etree.SubElement(te, f'{{{P}}}dissolve')
    else:etree.SubElement(te, f'{{{P}}}fade')
    
    # timing
    spids = [str(s.shape_id) for s in slide.shapes]
    timing = make_timing(spids)
    sld.append(timing)
    if clr is not None: sld.append(clr)

# ============ 内容（和之前一样） ============
s1=prs.slides.add_slide(prs.slide_layouts[6]);bg(s1,LIGHT)
rct(s1,Inches(1),Inches(2.0),Inches(0.06),Inches(2.5),BLUE)
tx(s1,Inches(1.5),Inches(2.0),Inches(10),Inches(1.0),"EasyApplyResume",sz=52,color=DARK,b=True,font='Trebuchet MS')
tx(s1,Inches(1.5),Inches(3.2),Inches(10),Inches(0.7),"AI-Driven Smart Job Platform",sz=24,color=BLUE,font='Trebuchet MS')
rct(s1,Inches(1.5),Inches(4.2),Inches(5),Inches(0.02),TEAL)
tx(s1,Inches(1.5),Inches(4.5),Inches(10),Inches(0.5),"Spring Boot . Spring AI . MyBatis-Plus . Redis . PostgreSQL . React . Docker",sz=13,color=GRAY)

s2=prs.slides.add_slide(prs.slide_layouts[6]);bg(s2,LIGHT);tbar(s2,"Problem Background");pn(s2,2)
for i,(ti,de,ac) in enumerate([("Resume Making Inefficient","Mass templates hard to choose\nManual formatting time-consuming",BLUE),("Info Asymmetry","Unfamiliar with enterprise needs\nLow JD-resume match rate",TEAL),("High Operation Burden","Mass data manual management\nLack of intelligent tools",ORANGE),("Lack Personalization","Generic templates not industry-fit\nMissing AI optimization",PURPLE)]):
    x=Inches(0.6+(i%2)*6.3);y=Inches(1.2+(i//2)*2.85)
    rctr(s2,x,y,Inches(5.8),Inches(2.4),CARD);rct(s2,x,y,Inches(0.06),Inches(2.4),ac)
    tx(s2,x+Inches(0.3),y+Inches(0.2),Inches(5.2),Inches(0.45),ti,sz=18,color=ac,b=True,font='Trebuchet MS')
    tx(s2,x+Inches(0.3),y+Inches(0.85),Inches(5.2),Inches(1.3),de,sz=14,color=TEXT)

s3=prs.slides.add_slide(prs.slide_layouts[6]);bg(s3,LIGHT);tbar(s3,"Tech Stack");pn(s3,3)
for ci,(cat,items,ac) in enumerate([("Foundation",["Spring Boot 3.3.5 + Java 21","Spring Security Dual Chain","MyBatis-Plus 3.5.7","Nacos Config Center"],BLUE),("AI & LLM",["Spring AI Multi-Model","ReAct Agent Architecture","RAG Retrieval Augmented","DashScope / Zhipu / Ollama"],TEAL),("Data & Storage",["MySQL + PostgreSQL / PgVector","Redis Cache & Messaging","MinIO / Qiniu OSS","Docker Deployment"],ORANGE)]):
    x=Inches(0.6+ci*4.1);rct(s3,x,Inches(1.2),Inches(3.8),Inches(0.5),ac)
    tx(s3,x+Inches(0.1),Inches(1.2),Inches(3.6),Inches(0.5),cat,sz=15,color=WHITE,b=True,font='Trebuchet MS',al=PP_ALIGN.CENTER,va=MSO_ANCHOR.MIDDLE)
    for j,item in enumerate(items):
        iy=Inches(2.0+j*0.55);rctr(s3,x,iy,Inches(3.8),Inches(0.45),CARD)
        tx(s3,x+Inches(0.15),iy+Inches(0.08),Inches(3.5),Inches(0.3),item,sz=11,color=TEXT,va=MSO_ANCHOR.MIDDLE)

s4=prs.slides.add_slide(prs.slide_layouts[6]);bg(s4,LIGHT);tbar(s4,"AI Agent Architecture");pn(s4,4)
for i,(nm,ds,ac) in enumerate([("BaseAgent","LLM Call . Memory . Base Abstraction",BLUE),("ReActAgent","Think -> Act -> Observe Cycle",TEAL),("ToolCallAgent","Tool Registry . Function Calling",ORANGE)]):
    y=Inches(1.2+i*1.5);rctr(s4,Inches(0.6),y,Inches(5.5),Inches(1.2),CARD);rct(s4,Inches(0.6),y,Inches(0.06),Inches(1.2),ac)
    tx(s4,Inches(1.0),y+Inches(0.15),Inches(5),Inches(0.4),nm,sz=20,color=ac,b=True,font='Trebuchet MS')
    tx(s4,Inches(1.0),y+Inches(0.6),Inches(5),Inches(0.5),ds,sz=12,color=GRAY)
    if i<2:tx(s4,Inches(3),y+Inches(1.15),Inches(1),Inches(0.3),"v",sz=16,color=GRAY,al=PP_ALIGN.CENTER)
tx(s4,Inches(6.8),Inches(1.2),Inches(6),Inches(0.5),"Core Components",sz=18,color=DARK,b=True,font='Trebuchet MS')
for i,(lb,ds,ac) in enumerate([("RAG Knowledge Base","PgVector + Cloud KB",BLUE),("Tool Calling","Web Search . Terminal . File . PDF",TEAL),("Chat Memory","MySQL Persist + In-Memory",ORANGE),("Security Filter","Sensitive Words . Auth",PURPLE),("MCP Client","Multi-Protocol Integration",BLUE),("Stream Output","SSE Real-Time Response",TEAL)]):
    y=Inches(1.85+i*0.82);rctr(s4,Inches(6.8),y,Inches(6),Inches(0.7),CARD);rct(s4,Inches(6.8),y,Inches(0.06),Inches(0.7),ac)
    tx(s4,Inches(7.1),y+Inches(0.05),Inches(5.4),Inches(0.3),lb,sz=14,color=ac,b=True,font='Trebuchet MS')
    tx(s4,Inches(7.1),y+Inches(0.38),Inches(5.4),Inches(0.3),ds,sz=11,color=GRAY)

s5=prs.slides.add_slide(prs.slide_layouts[6]);bg(s5,LIGHT);tbar(s5,"C-End . Resume AI Assistant");pn(s5,5)
for i,(ti,de,ac) in enumerate([("Smart Polish","AI analyzes & optimizes\nAuto adjust wording & layout",BLUE),("Custom Advice","Target JD-based\nSpecific modification plan",TEAL),("RAG Enhanced","Vector DB matching\nPrecise industry knowledge",ORANGE),("Tool Calling","Web Search . PDF . Email . File",PURPLE),("Stream Chat","SSE real-time response\nTypewriter effect",BLUE),("Eco Services","MinIO/Qiniu storage\nMulti-industry templates",TEAL)]):
    x=Inches(0.6+(i%3)*4.15);y=Inches(1.2+(i//3)*2.85)
    rctr(s5,x,y,Inches(3.8),Inches(2.45),CARD);rct(s5,x,y,Inches(3.8),Inches(0.05),ac)
    tx(s5,x+Inches(0.2),y+Inches(0.25),Inches(3.4),Inches(0.45),ti,sz=16,color=ac,b=True,font='Trebuchet MS')
    tx(s5,x+Inches(0.2),y+Inches(0.9),Inches(3.4),Inches(1.3),de,sz=12,color=TEXT)

s6=prs.slides.add_slide(prs.slide_layouts[6]);bg(s6,LIGHT);tbar(s6,"B-End . System Manager");pn(s6,6)
for i,(ti,de) in enumerate([("User Management","Register/Login . Permission . Role"),("Resume Template","Template CRUD . Enable/Disable"),("Written Test","Question CRUD . Category . Level"),("FAQ Management","Question maintenance . Review"),("User Guide","Manual editing . Rich text")]):
    y=Inches(1.15+i*1.12);rctr(s6,Inches(0.6),y,Inches(6.5),Inches(0.95),CARD);rct(s6,Inches(0.6),y,Inches(0.06),Inches(0.95),BLUE)
    tx(s6,Inches(1.0),y+Inches(0.1),Inches(5.8),Inches(0.3),ti,sz=16,color=BLUE,b=True,font='Trebuchet MS')
    tx(s6,Inches(1.0),y+Inches(0.45),Inches(5.8),Inches(0.4),de,sz=11,color=GRAY)
tx(s6,Inches(7.6),Inches(1.15),Inches(5.5),Inches(0.45),"AI Management Capabilities",sz=18,color=TEAL,b=True,font='Trebuchet MS')
for i,item in enumerate(["Smart Data Statistics & Analysis","Automated Operation Suggestions","Anomaly Detection & Alert","Smart Q&A Assistant","Content Auto Review","Batch Operation Automation","Log Analysis & Diagnosis"]):
    tx(s6,Inches(7.6),Inches(1.8+i*0.5),Inches(5.5),Inches(0.4),f"+ {item}",sz=13,color=TEXT)

s7=prs.slides.add_slide(prs.slide_layouts[6]);bg(s7,LIGHT);tbar(s7,"D-End . Data & Monitoring");pn(s7,7)
for i,(ti,de,ac) in enumerate([("Real-time Monitoring","DAU/Total PV tracking\nMulti-dim dashboard",BLUE),("Ad Management","Ad scheduling & control\nPerformance tracking",TEAL),("Server Monitor","SSH remote management\nHealth check & alert",ORANGE),("Smart Reports","Prometheus metrics\nGrafana dashboard",PURPLE)]):
    x=Inches(0.6+(i%2)*6.3);y=Inches(1.2+(i//2)*2.9)
    rctr(s7,x,y,Inches(5.8),Inches(2.5),CARD);rct(s7,x,y,Inches(5.8),Inches(0.05),ac)
    tx(s7,x+Inches(0.3),y+Inches(0.25),Inches(5.2),Inches(0.5),ti,sz=20,color=ac,b=True,font='Trebuchet MS')
    tx(s7,x+Inches(0.3),y+Inches(1.0),Inches(5.2),Inches(1.3),de,sz=14,color=TEXT)

s8=prs.slides.add_slide(prs.slide_layouts[6]);bg(s8,LIGHT);tbar(s8,"Future Prospects");pn(s8,8)
for ti,de,ac,x in [("Short-term","Improve multi-Agent\nOptimize RAG effect\nExpand template library",BLUE,Inches(1.0)),("Mid-term","Mobile app launch\nMore LLM integration\nOpen API ecosystem",TEAL,Inches(4.7)),("Long-term","Full-chain AI job platform\nFrom resume to onboarding\nIntelligent closed loop",ORANGE,Inches(8.4))]:
    rctr(s8,x,Inches(1.3),Inches(3.8),Inches(4.5),CARD);rct(s8,x,Inches(1.3),Inches(3.8),Inches(0.05),ac)
    tx(s8,x+Inches(0.3),Inches(1.6),Inches(3.2),Inches(0.5),ti,sz=22,color=ac,b=True,font='Trebuchet MS',al=PP_ALIGN.CENTER)
    tx(s8,x+Inches(0.3),Inches(2.4),Inches(3.2),Inches(3),de,sz=16,color=TEXT,al=PP_ALIGN.CENTER)
rct(s8,Inches(1),Inches(6.4),Inches(11.333),Inches(0.5),BLUE)
tx(s8,Inches(1),Inches(6.4),Inches(11.333),Inches(0.5),"Let everyone find their ideal job -- AI empowers the full recruitment process",sz=16,color=WHITE,b=True,font='Trebuchet MS',al=PP_ALIGN.CENTER,va=MSO_ANCHOR.MIDDLE)

s9=prs.slides.add_slide(prs.slide_layouts[6]);bg(s9,LIGHT);pn(s9,9)
rct(s9,Inches(5.5),Inches(2.0),Inches(2.333),Inches(0.04),BLUE)
tx(s9,Inches(1),Inches(2.3),Inches(11.333),Inches(1),"Thank You",sz=52,color=DARK,b=True,font='Trebuchet MS',al=PP_ALIGN.CENTER)
tx(s9,Inches(1),Inches(3.6),Inches(11.333),Inches(0.6),"EasyApplyResume -- AI-powered Smart Job & Management Platform",sz=18,color=BLUE,font='Trebuchet MS',al=PP_ALIGN.CENTER)
tx(s9,Inches(1),Inches(5.0),Inches(11.333),Inches(0.4),"Spring Boot . Spring AI . MyBatis-Plus . Redis . PostgreSQL . React . Docker",sz=12,color=GRAY,al=PP_ALIGN.CENTER)
tx(s9,Inches(1),Inches(5.6),Inches(11.333),Inches(0.4),"shiningCloud2025 2026",sz=11,color=GRAY,al=PP_ALIGN.CENTER)

# ============ 注入动画 ============
for slide in prs.slides:
    inject_all(slide, 'fade')

out="out/EasyApplyResume-项目介绍-v10.pptx"
prs.save(out)

# 验证
with zipfile.ZipFile(out,'r') as z:
    for f in z.namelist():
        parts=f.split('/')
        if len(parts)==3 and parts[0]=='ppt' and parts[1]=='slides':
            fn=parts[2]
            if fn.startswith('slide') and fn.endswith('.xml'):
                try:int(fn[5:-4])
                except:continue
                with z.open(f) as ff:
                    sld=etree.fromstring(ff.read())
                    order=[c.tag.split('}')[-1] for c in sld]
                    tm=sld.find(f'{{{P}}}timing')
                    ck=len(tm.findall(f'.//{{{P}}}cTn[@nodeType="clickEffect"]')) if tm else 0
                    print(f"{fn}: order=...{order[-3:]}, clickEffects={ck}")

print(f"\nSaved: {out}")
