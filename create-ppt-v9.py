"""
混合方案: python-pptx 做视觉 + ZIP 层面直接注入动画 XML
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree
import zipfile, io, uuid, os

P_NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'

etree.register_namespace('a', A_NS)
etree.register_namespace('r', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships')
etree.register_namespace('p', P_NS)

BG     = RGBColor(0xF8, 0xFA, 0xFC)
CARD   = RGBColor(0xFF, 0xFF, 0xFF)
BLUE   = RGBColor(0x02, 0x84, 0xC7)
TEAL   = RGBColor(0x0D, 0x94, 0x8F)
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
PURPLE = RGBColor(0x7C, 0x3A, 0xED)
DARK   = RGBColor(0x0F, 0x17, 0x2A)
TEXT   = RGBColor(0x1E, 0x29, 0x3B)
GRAY   = RGBColor(0x64, 0x74, 0x8B)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW = prs.slide_width; SH = prs.slide_height

def bg(s,c): s.background.fill.solid(); s.background.fill.fore_color.rgb = c
def rect(s,l,t,w,h,fl,li=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,l,t,w,h)
    sh.fill.solid();sh.fill.fore_color.rgb=fl
    if li:sh.line.fill.solid();sh.line.fill.fore_color.rgb=li
    else:sh.line.fill.background()
    return sh
def rrect(s,l,t,w,h,fl):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,l,t,w,h)
    sh.fill.solid();sh.fill.fore_color.rgb=fl;sh.line.fill.background()
    return sh
def txt(s,l,t,w,h,text,sz=12,color=TEXT,bold=False,font='Calibri',al=PP_ALIGN.LEFT,va=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(l,t,w,h);tf=tb.text_frame;tf.word_wrap=True
    p=tf.paragraphs[0];p.text=text;p.font.size=Pt(sz)
    p.font.color.rgb=color;p.font.bold=bold;p.font.name=font;p.alignment=al
    return tb
def tbar(s,title):
    rect(s,Inches(0),Inches(0),SW,Inches(0.8),DARK)
    txt(s,Inches(0.6),Inches(0.15),Inches(12),Inches(0.5),title,sz=22,color=WHITE,bold=True,font='Trebuchet MS',va=MSO_ANCHOR.MIDDLE)
    rect(s,Inches(0),Inches(0.8),SW,Inches(0.03),BLUE)
def pn(s,n):
    txt(s,Inches(11.5),Inches(7.1),Inches(1.5),Inches(0.3),f"{n}/9",sz=9,color=GRAY,al=PP_ALIGN.RIGHT)

# ============ 内容（不注入 python-pptx 层动画） ============
slides = []

# s1
s1=prs.slides.add_slide(prs.slide_layouts[6]);bg(s1,BG)
rect(s1,Inches(1),Inches(2.0),Inches(0.06),Inches(2.5),BLUE)
txt(s1,Inches(1.5),Inches(2.0),Inches(10),Inches(1.0),"EasyApplyResume",sz=52,color=DARK,bold=True,font='Trebuchet MS')
txt(s1,Inches(1.5),Inches(3.2),Inches(10),Inches(0.7),"AI 驱动的智能求职与管理平台",sz=24,color=BLUE,font='Trebuchet MS')
rect(s1,Inches(1.5),Inches(4.2),Inches(5),Inches(0.02),TEAL)
txt(s1,Inches(1.5),Inches(4.5),Inches(10),Inches(0.5),"Spring Boot . Spring AI . MyBatis-Plus . Redis . PostgreSQL . React . Docker",sz=13,color=GRAY)
slides.append(('s1','fade'))

# s2
s2=prs.slides.add_slide(prs.slide_layouts[6]);bg(s2,BG);tbar(s2,"Problem Background");pn(s2,2)
for i,(ti,de,ac) in enumerate([("Resume Making Inefficient","Mass templates hard to choose\nManual formatting time-consuming",BLUE),("Info Asymmetry","Unfamiliar with enterprise needs\nLow JD-resume match rate",TEAL),("High Operation Burden","Mass data manual management\nLack of intelligent tools",ORANGE),("Lack Personalization","Generic templates not industry-fit\nMissing AI optimization",PURPLE)]):
    x=Inches(0.6+(i%2)*6.3);y=Inches(1.2+(i//2)*2.85)
    rrect(s2,x,y,Inches(5.8),Inches(2.4),CARD)
    rect(s2,x,y,Inches(0.06),Inches(2.4),ac)
    txt(s2,x+Inches(0.3),y+Inches(0.2),Inches(5.2),Inches(0.45),ti,sz=18,color=ac,bold=True,font='Trebuchet MS')
    txt(s2,x+Inches(0.3),y+Inches(0.85),Inches(5.2),Inches(1.3),de,sz=14,color=TEXT)
slides.append(('s2','push'))

# s3
s3=prs.slides.add_slide(prs.slide_layouts[6]);bg(s3,BG);tbar(s3,"Tech Stack");pn(s3,3)
for ci,(cat,items,ac) in enumerate([("Foundation",["Spring Boot 3.3.5 + Java 21","Spring Security Dual Chain","MyBatis-Plus 3.5.7","Nacos Config Center"],BLUE),("AI & LLM",["Spring AI Multi-Model","ReAct Agent Architecture","RAG Retrieval Augmented","DashScope / Zhipu / Ollama"],TEAL),("Data & Storage",["MySQL + PostgreSQL / PgVector","Redis Cache & Messaging","MinIO / Qiniu OSS","Docker Deployment"],ORANGE)]):
    x=Inches(0.6+ci*4.1);rect(s3,x,Inches(1.2),Inches(3.8),Inches(0.5),ac)
    txt(s3,x+Inches(0.1),Inches(1.2),Inches(3.6),Inches(0.5),cat,sz=15,color=WHITE,bold=True,font='Trebuchet MS',al=PP_ALIGN.CENTER,va=MSO_ANCHOR.MIDDLE)
    for j,item in enumerate(items):
        iy=Inches(2.0+j*0.55);rrect(s3,x,iy,Inches(3.8),Inches(0.45),CARD)
        txt(s3,x+Inches(0.15),iy+Inches(0.08),Inches(3.5),Inches(0.3),item,sz=11,color=TEXT,va=MSO_ANCHOR.MIDDLE)
slides.append(('s3','cover'))

# s4
s4=prs.slides.add_slide(prs.slide_layouts[6]);bg(s4,BG);tbar(s4,"AI Agent Architecture");pn(s4,4)
for i,(nm,ds,ac) in enumerate([("BaseAgent","LLM Call . Memory . Base Abstraction",BLUE),("ReActAgent","Think -> Act -> Observe Cycle",TEAL),("ToolCallAgent","Tool Registry . Function Calling",ORANGE)]):
    y=Inches(1.2+i*1.5);rrect(s4,Inches(0.6),y,Inches(5.5),Inches(1.2),CARD)
    rect(s4,Inches(0.6),y,Inches(0.06),Inches(1.2),ac)
    txt(s4,Inches(1.0),y+Inches(0.15),Inches(5),Inches(0.4),nm,sz=20,color=ac,bold=True,font='Trebuchet MS')
    txt(s4,Inches(1.0),y+Inches(0.6),Inches(5),Inches(0.5),ds,sz=12,color=GRAY)
    if i<2:txt(s4,Inches(3),y+Inches(1.15),Inches(1),Inches(0.3),"v",sz=16,color=GRAY,al=PP_ALIGN.CENTER)
txt(s4,Inches(6.8),Inches(1.2),Inches(6),Inches(0.5),"Core Components",sz=18,color=DARK,bold=True,font='Trebuchet MS')
for i,(lb,ds,ac) in enumerate([("RAG Knowledge Base","PgVector + Cloud KB",BLUE),("Tool Calling","Web Search . Terminal . File . PDF",TEAL),("Chat Memory","MySQL Persist + In-Memory",ORANGE),("Security Filter","Sensitive Words . Auth",PURPLE),("MCP Client","Multi-Protocol Integration",BLUE),("Stream Output","SSE Real-Time Response",TEAL)]):
    y=Inches(1.85+i*0.82);rrect(s4,Inches(6.8),y,Inches(6),Inches(0.7),CARD)
    rect(s4,Inches(6.8),y,Inches(0.06),Inches(0.7),ac)
    txt(s4,Inches(7.1),y+Inches(0.05),Inches(5.4),Inches(0.3),lb,sz=14,color=ac,bold=True,font='Trebuchet MS')
    txt(s4,Inches(7.1),y+Inches(0.38),Inches(5.4),Inches(0.3),ds,sz=11,color=GRAY)
slides.append(('s4','wipe'))

# s5
s5=prs.slides.add_slide(prs.slide_layouts[6]);bg(s5,BG);tbar(s5,"C-End . Resume AI Assistant");pn(s5,5)
for i,(ti,de,ac) in enumerate([("Smart Polish","AI analyzes & optimizes\nAuto adjust wording & layout",BLUE),("Custom Advice","Target JD-based\nSpecific modification plan",TEAL),("RAG Enhanced","Vector DB matching\nPrecise industry knowledge",ORANGE),("Tool Calling","Web Search . PDF . Email . File",PURPLE),("Stream Chat","SSE real-time response\nTypewriter effect",BLUE),("Eco Services","MinIO/Qiniu storage\nMulti-industry templates",TEAL)]):
    x=Inches(0.6+(i%3)*4.15);y=Inches(1.2+(i//3)*2.85)
    rrect(s5,x,y,Inches(3.8),Inches(2.45),CARD)
    rect(s5,x,y,Inches(3.8),Inches(0.05),ac)
    txt(s5,x+Inches(0.2),y+Inches(0.25),Inches(3.4),Inches(0.45),ti,sz=16,color=ac,bold=True,font='Trebuchet MS')
    txt(s5,x+Inches(0.2),y+Inches(0.9),Inches(3.4),Inches(1.3),de,sz=12,color=TEXT)
slides.append(('s5','uncover'))

# s6
s6=prs.slides.add_slide(prs.slide_layouts[6]);bg(s6,BG);tbar(s6,"B-End . System Manager Assistant");pn(s6,6)
for i,(ti,de) in enumerate([("User Management","Register/Login . Permission . Role"),("Resume Template","Template CRUD . Enable/Disable"),("Written Test","Question CRUD . Category . Level"),("FAQ Management","Question maintenance . Review"),("User Guide","Manual editing . Rich text")]):
    y=Inches(1.15+i*1.12);rrect(s6,Inches(0.6),y,Inches(6.5),Inches(0.95),CARD)
    rect(s6,Inches(0.6),y,Inches(0.06),Inches(0.95),BLUE)
    txt(s6,Inches(1.0),y+Inches(0.1),Inches(5.8),Inches(0.3),ti,sz=16,color=BLUE,bold=True,font='Trebuchet MS')
    txt(s6,Inches(1.0),y+Inches(0.45),Inches(5.8),Inches(0.4),de,sz=11,color=GRAY)
txt(s6,Inches(7.6),Inches(1.15),Inches(5.5),Inches(0.45),"AI Management Capabilities",sz=18,color=TEAL,bold=True,font='Trebuchet MS')
for i,item in enumerate(["Smart Data Statistics & Analysis","Automated Operation Suggestions","Anomaly Detection & Alert","Smart Q&A Assistant","Content Auto Review","Batch Operation Automation","Log Analysis & Diagnosis"]):
    txt(s6,Inches(7.6),Inches(1.8+i*0.5),Inches(5.5),Inches(0.4),f"+ {item}",sz=13,color=TEXT)
slides.append(('s6','fade'))

# s7
s7=prs.slides.add_slide(prs.slide_layouts[6]);bg(s7,BG);tbar(s7,"D-End . Data & Monitoring Platform");pn(s7,7)
for i,(ti,de,ac) in enumerate([("Real-time Monitoring","DAU/Total PV tracking\nMulti-dim dashboard",BLUE),("Ad Management","Ad scheduling & control\nPerformance tracking",TEAL),("Server Monitor","SSH remote management\nHealth check & alert",ORANGE),("Smart Reports","Prometheus metrics\nGrafana dashboard",PURPLE)]):
    x=Inches(0.6+(i%2)*6.3);y=Inches(1.2+(i//2)*2.9)
    rrect(s7,x,y,Inches(5.8),Inches(2.5),CARD)
    rect(s7,x,y,Inches(5.8),Inches(0.05),ac)
    txt(s7,x+Inches(0.3),y+Inches(0.25),Inches(5.2),Inches(0.5),ti,sz=20,color=ac,bold=True,font='Trebuchet MS')
    txt(s7,x+Inches(0.3),y+Inches(1.0),Inches(5.2),Inches(1.3),de,sz=14,color=TEXT)
slides.append(('s7','push'))

# s8
s8=prs.slides.add_slide(prs.slide_layouts[6]);bg(s8,BG);tbar(s8,"Future Prospects");pn(s8,8)
for ti,de,ac,x in [("Short-term","Improve multi-Agent\nOptimize RAG effect\nExpand template library",BLUE,Inches(1.0)),("Mid-term","Mobile app launch\nMore LLM integration\nOpen API ecosystem",TEAL,Inches(4.7)),("Long-term","Full-chain AI job platform\nFrom resume to onboarding\nIntelligent closed loop",ORANGE,Inches(8.4))]:
    rrect(s8,x,Inches(1.3),Inches(3.8),Inches(4.5),CARD)
    rect(s8,x,Inches(1.3),Inches(3.8),Inches(0.05),ac)
    txt(s8,x+Inches(0.3),Inches(1.6),Inches(3.2),Inches(0.5),ti,sz=22,color=ac,bold=True,font='Trebuchet MS',al=PP_ALIGN.CENTER)
    txt(s8,x+Inches(0.3),Inches(2.4),Inches(3.2),Inches(3),de,sz=16,color=TEXT,al=PP_ALIGN.CENTER)
rect(s8,Inches(1),Inches(6.4),Inches(11.333),Inches(0.5),BLUE)
txt(s8,Inches(1),Inches(6.4),Inches(11.333),Inches(0.5),"Let everyone find their ideal job -- AI empowers the full recruitment process",sz=16,color=WHITE,bold=True,font='Trebuchet MS',al=PP_ALIGN.CENTER,va=MSO_ANCHOR.MIDDLE)
slides.append(('s8','cover'))

# s9
s9=prs.slides.add_slide(prs.slide_layouts[6]);bg(s9,BG);pn(s9,9)
rect(s9,Inches(5.5),Inches(2.0),Inches(2.333),Inches(0.04),BLUE)
txt(s9,Inches(1),Inches(2.3),Inches(11.333),Inches(1),"Thank You",sz=52,color=DARK,bold=True,font='Trebuchet MS',al=PP_ALIGN.CENTER)
txt(s9,Inches(1),Inches(3.6),Inches(11.333),Inches(0.6),"EasyApplyResume -- AI-powered Smart Job & Management Platform",sz=18,color=BLUE,font='Trebuchet MS',al=PP_ALIGN.CENTER)
txt(s9,Inches(1),Inches(5.0),Inches(11.333),Inches(0.4),"Spring Boot . Spring AI . MyBatis-Plus . Redis . PostgreSQL . React . Docker",sz=12,color=GRAY,al=PP_ALIGN.CENTER)
txt(s9,Inches(1),Inches(5.6),Inches(11.333),Inches(0.4),"shiningCloud2025 2026",sz=11,color=GRAY,al=PP_ALIGN.CENTER)
slides.append(('s9','dissolve'))

# ============ 保存为基础文件 ============
base_path = "out/EasyApplyResume-项目介绍-v9-base.pptx"
prs.save(base_path)

# ============ ZIP 层面注入动画 XML ============
tmap = {'push':'l','cover':'r','wipe':'l','uncover':'d'}
out_path = "out/EasyApplyResume-项目介绍-v9.pptx"

with zipfile.ZipFile(base_path, 'r') as zin:
    with zipfile.ZipFile(out_path, 'w', zipfile.ZIP_DEFLATED) as zout:
        slide_files = []
        for f in zin.namelist():
            parts = f.split('/')
            if len(parts) == 3 and parts[0] == 'ppt' and parts[1] == 'slides':
                fn = parts[2]
                if fn.startswith('slide') and fn.endswith('.xml'):
                    num_str = fn[5:-4]
                    try:
                        int(num_str)
                        slide_files.append(f)
                    except: pass
        slide_files.sort(key=lambda x: int(x.split('/')[-1][5:-4]))
        
        for fname in zin.infolist():
            if fname.filename in slide_files:
                idx = slide_files.index(fname.filename)
                trans = slides[idx][1]
                
                xml_bytes = zin.read(fname.filename)
                sld = etree.fromstring(xml_bytes)
                
                # 移除 clrMapOvr
                clr = sld.find(f'{{{P_NS}}}clrMapOvr')
                if clr is not None: sld.remove(clr)
                
                # 过渡
                trans_el = etree.Element(f'{{{P_NS}}}transition', spd='med')
                if trans in tmap:
                    t = etree.SubElement(trans_el, f'{{{P_NS}}}{trans}')
                    t.set('dir', tmap[trans])
                elif trans == 'dissolve':
                    etree.SubElement(trans_el, f'{{{P_NS}}}dissolve')
                else:
                    etree.SubElement(trans_el, f'{{{P_NS}}}fade')
                sld.append(trans_el)
                
                # 动画：收集所有 shape id
                cSld = sld.find(f'{{{P_NS}}}cSld')
                spTree = cSld.find(f'{{{P_NS}}}spTree')
                shape_ids = []
                for sp in spTree.findall(f'{{{P_NS}}}sp'):
                    nv = sp.find(f'{{{P_NS}}}nvSpPr')
                    if nv is not None:
                        cnv = nv.find(f'{{{P_NS}}}cNvPr')
                        if cnv is not None and cnv.get('id'):
                            shape_ids.append(cnv.get('id'))
                # 也收集 grpSp 和 pic
                for tag in ('grpSp','pic','graphicFrame'):
                    for el in spTree.findall(f'{{{P_NS}}}{tag}'):
                        nv = el.find(f'{{{P_NS}}}nvSpPr') or el.find(f'{{{P_NS}}}nvGrpSpPr') or el.find(f'{{{P_NS}}}nvPicPr') or el.find(f'{{{P_NS}}}nvGraphicFramePr')
                        if nv is not None:
                            cnv = nv.find(f'{{{P_NS}}}cNvPr')
                            if cnv is not None and cnv.get('id'):
                                shape_ids.append(cnv.get('id'))
                
                # 构建 timing
                timing = etree.Element(f'{{{P_NS}}}timing')
                tnLst = etree.SubElement(timing, f'{{{P_NS}}}tnLst')
                seq = etree.SubElement(tnLst, f'{{{P_NS}}}seq', concurrent='0')
                ct_root = etree.SubElement(seq, f'{{{P_NS}}}cTn', id=str(uuid.uuid4()), dur='indefinite')
                scl = etree.SubElement(ct_root, f'{{{P_NS}}}stCondLst')
                etree.SubElement(scl, f'{{{P_NS}}}cond', delay='0')
                childTnLst = etree.SubElement(seq, f'{{{P_NS}}}childTnLst')
                
                for spid in shape_ids:
                    par = etree.SubElement(childTnLst, f'{{{P_NS}}}par')
                    ct = etree.SubElement(par, f'{{{P_NS}}}cTn', id=str(uuid.uuid4()), fill='hold')
                    stc = etree.SubElement(ct, f'{{{P_NS}}}stCondLst')
                    cond = etree.SubElement(stc, f'{{{P_NS}}}cond', evt='onclick', delay='0')
                    te = etree.SubElement(cond, f'{{{P_NS}}}tgtEl')
                    etree.SubElement(te, f'{{{P_NS}}}sldTgt')
                    cl = etree.SubElement(par, f'{{{P_NS}}}childTnLst')
                    ae = etree.SubElement(cl, f'{{{P_NS}}}animEffect', transition='in', filter='appear')
                    cb = etree.SubElement(ae, f'{{{P_NS}}}cBhvr')
                    etree.SubElement(cb, f'{{{P_NS}}}cTn', id=str(uuid.uuid4()), dur='500')
                    tgt = etree.SubElement(cb, f'{{{P_NS}}}tgtEl')
                    etree.SubElement(tgt, f'{{{P_NS}}}spTgt', spid=spid)
                
                sld.append(timing)
                sld.append(clr) if clr is not None else None
                
                zout.writestr(fname, etree.tostring(sld, xml_declaration=True, encoding='UTF-8', standalone=True))
            else:
                zout.writestr(fname, zin.read(fname.filename))

# 验证
with zipfile.ZipFile(out_path, 'r') as z:
    for f in sorted(z.namelist()):
        parts = f.split('/')
        if len(parts) == 3 and parts[0] == 'ppt' and parts[1] == 'slides':
            fn = parts[2]
            if fn.startswith('slide') and fn.endswith('.xml'):
                num_str = fn[5:-4]
                try: int(num_str)
                except: continue
                with z.open(f) as ff:
                    sld = etree.fromstring(ff.read())
                    order = [c.tag.split('}')[-1] for c in sld]
                    timing = sld.find(f'{{{P_NS}}}timing')
                    pars = len(timing.findall(f'.//{{{P_NS}}}par')) if timing is not None else 0
                    ae = len(timing.findall(f'.//{{{P_NS}}}animEffect')) if timing is not None else 0
                    has_onclick = timing is not None and any(c.get('evt')=='onclick' for c in timing.findall(f'.//{{{P_NS}}}cond'))
                    print(f"{fn}: order=...{order[-3:]}, pars={pars}, ae={ae}, onclick={has_onclick}")

print(f"\nSaved: {out_path}")
