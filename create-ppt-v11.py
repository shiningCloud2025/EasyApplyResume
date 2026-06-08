"""
从 native-anim.pptx 提取精确的 timing XML 模板，复制到每页
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree
import zipfile, uuid, copy

P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

# ============ 从 native-anim.pptx 提取 timing 模板 ============
def extract_timing_template():
    with zipfile.ZipFile('out/native-anim.pptx', 'r') as z:
        with z.open('ppt/slides/slide1.xml') as f:
            sld = etree.fromstring(f.read())
    return sld.find(f'{{{P}}}timing')

TIMING_TEMPLATE = extract_timing_template()
print(f"Template extracted: {len(TIMING_TEMPLATE.findall(f'.//{{{P}}}par'))} pars")

def clone_timing(shape_ids):
    """深拷贝模板 timing，替换 shape spid"""
    timing = copy.deepcopy(TIMING_TEMPLATE)
    # 替换所有 spTgt 的 spid
    all_targets = timing.findall(f'.//{{{P}}}spTgt')
    for i, tgt in enumerate(all_targets):
        if i < len(shape_ids):
            tgt.set('spid', str(shape_ids[i]))
    # 移除多余 targets（原生有形状+文字的，我们只有形状）
    # 只保留带 clickEffect 的 par 对应数量的 target
    click_effects = timing.findall(f'.//{{{P}}}cTn[@nodeType="clickEffect"]')
    target_count = len(shape_ids)
    
    # 实际上我们需要重建 clickEffect 部分以精确匹配 shape_ids 数量
    # 找到 clickEffect 的父 p:par
    # 简化方案：找到所有 clickEffect cTn 的父 par，复制最后一个并修改 spid
    click_pars = []
    for ce in click_effects:
        parent = ce.getparent()
        while parent is not None and parent.tag != f'{{{P}}}par':
            parent = parent.getparent()
        if parent is not None:
            click_pars.append(parent)
    
    if len(click_pars) > len(shape_ids):
        # 移除多余的
        parent_of_pars = click_pars[0].getparent()
        for p in click_pars[len(shape_ids):]:
            parent_of_pars.remove(p)
    elif len(click_pars) < len(shape_ids):
        # 添加更多
        parent_of_pars = click_pars[0].getparent()
        template_par = click_pars[-1]
        for i in range(len(click_pars), len(shape_ids)):
            new_par = copy.deepcopy(template_par)
            parent_of_pars.append(new_par)
    
    # 更新所有 spTgt
    updated_targets = timing.findall(f'.//{{{P}}}spTgt')
    for i, tgt in enumerate(updated_targets):
        if i < len(shape_ids):
            tgt.set('spid', str(shape_ids[i]))
    
    return timing

# ============ PPT 内容 ============
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
SW=prs.slide_width;SH=prs.slide_height

BG=CARD=RGBColor(0xFF,0xFF,0xFF);BLUE=RGBColor(0x02,0x84,0xC7)
TEAL=RGBColor(0x0D,0x94,0x8F);ORANGE=RGBColor(0xEA,0x58,0x0C)
PURPLE=RGBColor(0x7C,0x3A,0xED);DARK=RGBColor(0x0F,0x17,0x2A)
TEXT=RGBColor(0x1E,0x29,0x3B);GRAY=RGBColor(0x64,0x74,0x8B)
WHITE=RGBColor(0xFF,0xFF,0xFF);LIGHT=RGBColor(0xF1,0xF5,0xF9)

def bg(s,c):s.background.fill.solid();s.background.fill.fore_color.rgb=c
def rct(s,l,t,w,h,fl,li=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,l,t,w,h)
    sh.fill.solid();sh.fill.fore_color.rgb=fl
    if li:sh.line.fill.solid();sh.line.fill.fore_color.rgb=li
    else:sh.line.fill.background()
    return sh
def rctr(s,l,t,w,h,fl):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,l,t,w,h)
    sh.fill.solid();sh.fill.fore_color.rgb=fl;sh.line.fill.background()
    return sh
def tx(s,l,t,w,h,text,sz=12,color=TEXT,b=False,font='Calibri',al=PP_ALIGN.LEFT,va=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(l,t,w,h);tf=tb.text_frame;tf.word_wrap=True
    p=tf.paragraphs[0];p.text=text;p.font.size=Pt(sz)
    p.font.color.rgb=color;p.font.bold=b;p.font.name=font;p.alignment=al
    return tb
def tbar(s,tt):
    rct(s,Inches(0),Inches(0),SW,Inches(0.8),DARK)
    tx(s,Inches(0.6),Inches(0.15),Inches(12),Inches(0.5),tt,sz=22,color=WHITE,b=True,font='Trebuchet MS',va=MSO_ANCHOR.MIDDLE)
    rct(s,Inches(0),Inches(0.8),SW,Inches(0.03),BLUE)
def pn(s,n):
    tx(s,Inches(11.5),Inches(7.1),Inches(1.5),Inches(0.3),f"{n}/9",sz=9,color=GRAY,al=PP_ALIGN.RIGHT)

# ============ 9 页内容 ============
s1=prs.slides.add_slide(prs.slide_layouts[6]);bg(s1,LIGHT)
rct(s1,Inches(1),Inches(2.0),Inches(0.06),Inches(2.5),BLUE)
tx(s1,Inches(1.5),Inches(2.0),Inches(10),Inches(1.0),"EasyApplyResume",sz=52,color=DARK,b=True,font='Trebuchet MS')
tx(s1,Inches(1.5),Inches(3.2),Inches(10),Inches(0.7),"AI-Driven Smart Job Platform",sz=24,color=BLUE,font='Trebuchet MS')
rct(s1,Inches(1.5),Inches(4.2),Inches(5),Inches(0.02),TEAL)
tx(s1,Inches(1.5),Inches(4.5),Inches(10),Inches(0.5),"Spring Boot . Spring AI . MyBatis-Plus . Redis . PostgreSQL . React . Docker",sz=13,color=GRAY)

s2=prs.slides.add_slide(prs.slide_layouts[6]);bg(s2,LIGHT);tbar(s2,"Problem Background");pn(s2,2)
for i,(ti,de,ac) in enumerate([("Resume Inefficient","Mass templates hard to choose\nManual formatting time-consuming",BLUE),("Info Asymmetry","Unfamiliar with enterprise needs\nLow JD-resume match rate",TEAL),("Operation Burden","Mass data manual management\nLack of intelligent tools",ORANGE),("Lack Personalization","Generic not industry-fit\nMissing AI optimization",PURPLE)]):
    x=Inches(0.6+(i%2)*6.3);y=Inches(1.2+(i//2)*2.85)
    rctr(s2,x,y,Inches(5.8),Inches(2.4),CARD);rct(s2,x,y,Inches(0.06),Inches(2.4),ac)
    tx(s2,x+Inches(0.3),y+Inches(0.2),Inches(5.2),Inches(0.45),ti,sz=18,color=ac,b=True,font='Trebuchet MS')
    tx(s2,x+Inches(0.3),y+Inches(0.85),Inches(5.2),Inches(1.3),de,sz=14,color=TEXT)

s3=prs.slides.add_slide(prs.slide_layouts[6]);bg(s3,LIGHT);tbar(s3,"Tech Stack");pn(s3,3)
for ci,(cat,items,ac) in enumerate([("Foundation",["Spring Boot 3.3.5 + Java 21","Spring Security Dual","MyBatis-Plus 3.5.7","Nacos Config Center"],BLUE),("AI & LLM",["Spring AI Multi-Model","ReAct Agent","RAG Retrieval","DashScope / Zhipu"],TEAL),("Data",["MySQL + PgVector","Redis Cache & Messaging","MinIO / Qiniu OSS","Docker Deployment"],ORANGE)]):
    x=Inches(0.6+ci*4.1);rct(s3,x,Inches(1.2),Inches(3.8),Inches(0.5),ac)
    tx(s3,x+Inches(0.1),Inches(1.2),Inches(3.6),Inches(0.5),cat,sz=15,color=WHITE,b=True,font='Trebuchet MS',al=PP_ALIGN.CENTER,va=MSO_ANCHOR.MIDDLE)
    for j,item in enumerate(items):
        iy=Inches(2.0+j*0.55);rctr(s3,x,iy,Inches(3.8),Inches(0.45),CARD)
        tx(s3,x+Inches(0.15),iy+Inches(0.08),Inches(3.5),Inches(0.3),item,sz=11,color=TEXT,va=MSO_ANCHOR.MIDDLE)

s4=prs.slides.add_slide(prs.slide_layouts[6]);bg(s4,LIGHT);tbar(s4,"AI Agent Architecture");pn(s4,4)
for i,(nm,ds,ac) in enumerate([("BaseAgent","LLM Call . Memory . Base",BLUE),("ReActAgent","Think -> Act -> Observe",TEAL),("ToolCallAgent","Tool Registry . Function Call",ORANGE)]):
    y=Inches(1.2+i*1.5);rctr(s4,Inches(0.6),y,Inches(5.5),Inches(1.2),CARD);rct(s4,Inches(0.6),y,Inches(0.06),Inches(1.2),ac)
    tx(s4,Inches(1.0),y+Inches(0.15),Inches(5),Inches(0.4),nm,sz=20,color=ac,b=True,font='Trebuchet MS')
    tx(s4,Inches(1.0),y+Inches(0.6),Inches(5),Inches(0.5),ds,sz=12,color=GRAY)
    if i<2:tx(s4,Inches(3),y+Inches(1.15),Inches(1),Inches(0.3),"v",sz=16,color=GRAY,al=PP_ALIGN.CENTER)
tx(s4,Inches(6.8),Inches(1.2),Inches(6),Inches(0.5),"Core Components",sz=18,color=DARK,b=True,font='Trebuchet MS')
for i,(lb,ds,ac) in enumerate([("RAG Knowledge","PgVector + Cloud KB",BLUE),("Tool Calling","Web Search . Terminal . PDF",TEAL),("Chat Memory","MySQL Persist + InMemory",ORANGE),("Security","Sensitive Words . Auth",PURPLE),("MCP Client","Multi-Protocol",BLUE),("SSE Stream","Real-Time Response",TEAL)]):
    y=Inches(1.85+i*0.82);rctr(s4,Inches(6.8),y,Inches(6),Inches(0.7),CARD);rct(s4,Inches(6.8),y,Inches(0.06),Inches(0.7),ac)
    tx(s4,Inches(7.1),y+Inches(0.05),Inches(5.4),Inches(0.3),lb,sz=14,color=ac,b=True,font='Trebuchet MS')
    tx(s4,Inches(7.1),y+Inches(0.38),Inches(5.4),Inches(0.3),ds,sz=11,color=GRAY)

s5=prs.slides.add_slide(prs.slide_layouts[6]);bg(s5,LIGHT);tbar(s5,"C-End . Resume AI Assistant");pn(s5,5)
for i,(ti,de,ac) in enumerate([("Smart Polish","AI analyzes & optimizes\nAuto adjust wording & layout",BLUE),("Custom Advice","Target JD-based\nSpecific modification",TEAL),("RAG Enhanced","Vector DB matching\nPrecise industry knowledge",ORANGE),("Tool Calling","Web Search . PDF . Email . File",PURPLE),("Stream Chat","SSE real-time response\nTypewriter effect",BLUE),("Eco Services","MinIO/Qiniu storage\nMulti-industry templates",TEAL)]):
    x=Inches(0.6+(i%3)*4.15);y=Inches(1.2+(i//3)*2.85)
    rctr(s5,x,y,Inches(3.8),Inches(2.45),CARD);rct(s5,x,y,Inches(3.8),Inches(0.05),ac)
    tx(s5,x+Inches(0.2),y+Inches(0.25),Inches(3.4),Inches(0.45),ti,sz=16,color=ac,b=True,font='Trebuchet MS')
    tx(s5,x+Inches(0.2),y+Inches(0.9),Inches(3.4),Inches(1.3),de,sz=12,color=TEXT)

s6=prs.slides.add_slide(prs.slide_layouts[6]);bg(s6,LIGHT);tbar(s6,"B-End . System Manager");pn(s6,6)
for i,(ti,de) in enumerate([("User Management","Register/Login . Permission"),("Resume Template","Template CRUD . Enable/Disable"),("Written Test","Question CRUD . Category . Level"),("FAQ Management","Question maintenance . Review"),("User Guide","Manual editing . Rich text")]):
    y=Inches(1.15+i*1.12);rctr(s6,Inches(0.6),y,Inches(6.5),Inches(0.95),CARD);rct(s6,Inches(0.6),y,Inches(0.06),Inches(0.95),BLUE)
    tx(s6,Inches(1.0),y+Inches(0.1),Inches(5.8),Inches(0.3),ti,sz=16,color=BLUE,b=True,font='Trebuchet MS')
    tx(s6,Inches(1.0),y+Inches(0.45),Inches(5.8),Inches(0.4),de,sz=11,color=GRAY)
tx(s6,Inches(7.6),Inches(1.15),Inches(5.5),Inches(0.45),"AI Management",sz=18,color=TEAL,b=True,font='Trebuchet MS')
for i,item in enumerate(["Smart Data Analysis","Operation Suggestions","Anomaly Detection","Smart Q&A Assistant","Content Auto Review","Batch Automation","Log Diagnosis"]):
    tx(s6,Inches(7.6),Inches(1.8+i*0.5),Inches(5.5),Inches(0.4),f"+ {item}",sz=13,color=TEXT)

s7=prs.slides.add_slide(prs.slide_layouts[6]);bg(s7,LIGHT);tbar(s7,"D-End . Data & Monitoring");pn(s7,7)
for i,(ti,de,ac) in enumerate([("Real-time Monitor","DAU/Total PV tracking\nMulti-dim dashboard",BLUE),("Ad Management","Ad scheduling & control\nPerformance tracking",TEAL),("Server Monitor","SSH remote management\nHealth check & alert",ORANGE),("Smart Reports","Prometheus metrics\nGrafana dashboard",PURPLE)]):
    x=Inches(0.6+(i%2)*6.3);y=Inches(1.2+(i//2)*2.9)
    rctr(s7,x,y,Inches(5.8),Inches(2.5),CARD);rct(s7,x,y,Inches(5.8),Inches(0.05),ac)
    tx(s7,x+Inches(0.3),y+Inches(0.25),Inches(5.2),Inches(0.5),ti,sz=20,color=ac,b=True,font='Trebuchet MS')
    tx(s7,x+Inches(0.3),y+Inches(1.0),Inches(5.2),Inches(1.3),de,sz=14,color=TEXT)

s8=prs.slides.add_slide(prs.slide_layouts[6]);bg(s8,LIGHT);tbar(s8,"Future Prospects");pn(s8,8)
for ti,de,ac,x in [("Short-term","Improve multi-Agent\nOptimize RAG effect\nExpand template library",BLUE,Inches(1.0)),("Mid-term","Mobile app launch\nMore LLM integration\nOpen API ecosystem",TEAL,Inches(4.7)),("Long-term","Full-chain AI job platform\nFrom resume to onboarding\nIntelligent closed loop",ORANGE,Inches(8.4))]:
    rctr(s8,x,Inches(1.3),Inches(3.8),Inches(4.5),CARD);rct(s8,x,Inches(1.3),Inches(3.8),Inches(0.05),ac)
    tx(s8,x+Inches(0.3),Inches(1.6),Inches(3.2),Inches(0.5),ti,sz=22,color=ac,b=True,font='Trebuchet MS',al=PP_ALIGN.CENTER)
    tx(s8,x+Inches(0.3),Inches(2.4),Inches(3.2),Inches(3),de,sz=16,color=TEXT,al=PP_ALIGN.CENTER)
rct(s8,Inches(1),Inches(6.4),Inches(11.333),Inches(0.5),BLUE)
tx(s8,Inches(1),Inches(6.4),Inches(11.333),Inches(0.5),"Let everyone find their ideal job",sz=16,color=WHITE,b=True,font='Trebuchet MS',al=PP_ALIGN.CENTER,va=MSO_ANCHOR.MIDDLE)

s9=prs.slides.add_slide(prs.slide_layouts[6]);bg(s9,LIGHT);pn(s9,9)
rct(s9,Inches(5.5),Inches(2.0),Inches(2.333),Inches(0.04),BLUE)
tx(s9,Inches(1),Inches(2.3),Inches(11.333),Inches(1),"Thank You",sz=52,color=DARK,b=True,font='Trebuchet MS',al=PP_ALIGN.CENTER)
tx(s9,Inches(1),Inches(3.6),Inches(11.333),Inches(0.6),"EasyApplyResume -- AI-powered Smart Job Platform",sz=18,color=BLUE,font='Trebuchet MS',al=PP_ALIGN.CENTER)
tx(s9,Inches(1),Inches(5.0),Inches(11.333),Inches(0.4),"Spring Boot . Spring AI . MyBatis-Plus . Redis . PostgreSQL . React . Docker",sz=12,color=GRAY,al=PP_ALIGN.CENTER)
tx(s9,Inches(1),Inches(5.6),Inches(11.333),Inches(0.4),"shiningCloud2025 2026",sz=11,color=GRAY,al=PP_ALIGN.CENTER)

# ============ 保存 ============
out="out/EasyApplyResume-项目介绍-v11.pptx"
prs.save(out)

# ============ ZIP 层面注入 timing ============
with zipfile.ZipFile(out,'r') as zin:
    with zipfile.ZipFile(out+'.tmp','w',zipfile.ZIP_DEFLATED) as zout:
        slide_files=[]
        for f in zin.namelist():
            parts=f.split('/')
            if len(parts)==3 and parts[0]=='ppt' and parts[1]=='slides':
                fn=parts[2]
                if fn.startswith('slide') and fn.endswith('.xml'):
                    try:int(fn[5:-4]);slide_files.append(f)
                    except:pass
        slide_files.sort(key=lambda x:int(x.split('/')[-1][5:-4]))
        
        for fname in zin.infolist():
            data=zin.read(fname.filename)
            if fname.filename in slide_files:
                sld=etree.fromstring(data)
                # 移除 transition 和 clrMapOvr
                for tag in ('transition','timing','clrMapOvr'):
                    el=sld.find(f'{{{P}}}{tag}')
                    if el is not None:sld.remove(el)
                
                # 收集 shape IDs
                cSld=sld.find(f'{{{P}}}cSld')
                spTree=cSld.find(f'{{{P}}}spTree')
                all_ids=[]
                for tag in ('sp','grpSp','pic','graphicFrame'):
                    for el in spTree.findall(f'{{{P}}}{tag}'):
                        nv=el.find(f'{{{P}}}nvSpPr') or el.find(f'{{{P}}}nvGrpSpPr') or el.find(f'{{{P}}}nvPicPr') or el.find(f'{{{P}}}nvGraphicFramePr')
                        if nv is not None:
                            cnv=nv.find(f'{{{P}}}cNvPr')
                            if cnv is not None and cnv.get('id'):
                                all_ids.append(int(cnv.get('id')))
                
                all_ids.sort()
                timing=clone_timing(all_ids)
                sld.append(timing)
                # clrMapOvr 放最后
                clr=etree.Element(f'{{{P}}}clrMapOvr')
                sld.append(clr)
                
                data=etree.tostring(sld,xml_declaration=True,encoding='UTF-8',standalone=True)
            zout.writestr(fname,data)

import os
os.replace(out+'.tmp',out)

# 验证
with zipfile.ZipFile(out,'r') as z:
    for f in sorted(z.namelist()):
        parts=f.split('/')
        if len(parts)==3 and parts[0]=='ppt' and parts[1]=='slides':
            fn=parts[2]
            if fn.startswith('slide') and fn.endswith('.xml'):
                try:int(fn[5:-4])
                except:continue
                with z.open(f) as ff:
                    sld=etree.fromstring(ff.read())
                    tm=sld.find(f'{{{P}}}timing')
                    ck=len(tm.findall(f'.//{{{P}}}cTn[@nodeType="clickEffect"]')) if tm is not None else 0
                    tr=sld.find(f'{{{P}}}transition')
                    print(f"{fn}: clickEffects={ck}, hasTransition={tr is not None}")

print(f"\nSaved: {out}")
